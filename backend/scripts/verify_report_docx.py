"""Validate the generated .docx: opens cleanly, structure and content intact."""
from __future__ import annotations

import pathlib
import zipfile

from docx import Document

ROOT = pathlib.Path(__file__).resolve().parent.parent.parent
DOC = ROOT / "docs" / "RBAC_AUDIT_AND_TESTING_REPORT.docx"

print(f"file      : {DOC.name}")
print(f"size      : {DOC.stat().st_size:,} bytes")

# valid OOXML package?
with zipfile.ZipFile(DOC) as z:
    bad = z.testzip()
    assert bad is None, f"corrupt member: {bad}"
    names = z.namelist()
assert "word/document.xml" in names, "missing word/document.xml"
assert "[Content_Types].xml" in names, "missing content types"
print(f"ooxml     : valid zip, {len(names)} parts")

doc = Document(DOC)
paras = doc.paragraphs
tables = doc.tables
print(f"paragraphs: {len(paras)}")
print(f"tables    : {len(tables)}")
print(f"sections  : {len(doc.sections)}")

headings = [p.text for p in paras if p.style.name.startswith("Heading") or
            p.style.name == "Title"]
print(f"headings  : {len(headings)}")
print()
for h in headings:
    indent = "  " if not h[:2].strip().isdigit() else ""
    print(f"{indent}{h}")

# content assertions
full = "\n".join(p.text for p in paras)
for t in tables:
    for row in t.rows:
        full += "\n" + " | ".join(c.text for c in row.cells)

required = {
    "DEBUG fail-closed": "DEBUG defaulted to True",
    "freeze 503": "FREEZE_STATE_UNAVAILABLE",
    "approver authority": "authority list let any approver",
    "assign validation": "ASSIGNEE_NOT_FOUND",
    "acquisition authority": "acquisition_flow.py",
    "audit hygiene": "except Exception: pass",
    "permission drift": "23 canonical permissions",
    "test counts": "723 passed",
    "baseline delta": "698 passed",
    "boot test": "assert_secure_secret",
    "7 identities": "system_admin",
    "error table": "SELF_DISABLE_DENIED",
    "remaining gap": "Persist scope assignments",
    "verification harness": "42 passed",
}
print()
missing = []
for label, needle in required.items():
    ok = needle in full
    if not ok:
        missing.append(label)
    print(f"  {'OK  ' if ok else 'MISS'} {label}")

print()
assert not missing, f"missing content: {missing}"
print("RESULT: document valid, all required content present")
